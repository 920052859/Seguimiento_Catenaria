"""Extract a defensible seven-activity OCS baseline to JSON.

Usage: python extract_schedule.py input.xlsx output.json
The source workbook is not copied to the public dashboard repository.
"""
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

DIRECT_MAPPING = {
    "Taladro de Bóveda de Túnel": ["perforacion"],
    "Montaje de Espárragos": ["fijacion"],
    "Montaje de Tubo + Estribo + Suspensión": ["soporte", "mensula", "aislador"],
    "Montaje de Perfil CR": ["perfil"],
    "Tendido de Hilo de Contacto": ["hilo"],
}
EXCLUDED_SCOPE = [
    "Marcación en Túnel",
    "Montaje de Aislador de Sección",
    "Montaje de Conexiones de Alimentación",
    "Montaje de Conexiones Eléctricas y Puesta a Tierra 1",
    "Montaje de Conexiones Eléctricas y Puesta a Tierra 2",
    "Montaje de Conexiones Eléctricas y Puesta a Tierra 3",
]


def iso(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def extract(path: Path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    stages = {}
    issues = []

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        detail_end = ws.max_row + 1
        for row in range(1, ws.max_row + 1):
            if str(ws.cell(row, 1).value or "").strip().upper() == "RESUMEN Y AJUSTE":
                detail_end = row
                break

        tunnels = {}
        current_tunnel = None
        for row in range(1, detail_end):
            label = str(ws.cell(row, 1).value or "").strip()
            tunnel_match = re.fullmatch(r"TUNEL\s+(\d+)", label, flags=re.IGNORECASE)
            if tunnel_match:
                current_tunnel = f"Tunel {int(tunnel_match.group(1))}"
                tunnels[current_tunnel] = {}
                continue
            if not current_tunnel or label not in DIRECT_MAPPING:
                continue
            start = iso(ws.cell(row, 3).value)
            end = iso(ws.cell(row, 4).value)
            if not start or not end:
                issues.append({"sheet": sheet_name, "row": row, "activity": label, "issue": "fecha detallada ausente"})
                continue
            if start > end:
                issues.append({"sheet": sheet_name, "row": row, "activity": label, "start": start, "end": end, "issue": "inicio posterior al fin"})
                continue
            for activity in DIRECT_MAPPING[label]:
                window = {"start": start, "end": end, "source_category": label}
                if activity == "aislador":
                    window["mapping_note"] = "proxy provisional: aislador estándar no desagregado en el cronograma"
                tunnels[current_tunnel][activity] = window

        stages[sheet_name] = {"tunnels": tunnels}

        for row in range(1, ws.max_row + 1):
            start = iso(ws.cell(row, 3).value)
            end = iso(ws.cell(row, 4).value)
            if start and end and start > end:
                issues.append({
                    "sheet": sheet_name,
                    "row": row,
                    "activity": str(ws.cell(row, 1).value or "").strip(),
                    "start": start,
                    "end": end,
                    "issue": "inicio posterior al fin en fuente",
                })

    expected = {"Tramo 3": 5, "Tramo 4": 4}
    for stage, tunnel_count in expected.items():
        actual = len(stages.get(stage, {}).get("tunnels", {}))
        if actual != tunnel_count:
            issues.append({"sheet": stage, "issue": f"se esperaban {tunnel_count} túneles y se extrajeron {actual}"})

    return {
        "schema_version": 2,
        "source_filename": path.name,
        "basis": "Cronograma detallado por túnel (filas previas a RESUMEN Y AJUSTE)",
        "methodology": {
            "perforacion": "Taladro de Bóveda de Túnel",
            "fijacion": "Montaje de Espárragos",
            "soporte": "Montaje de Tubo + Estribo + Suspensión",
            "mensula": "Montaje de Tubo + Estribo + Suspensión",
            "aislador": "proxy de Montaje de Tubo + Estribo + Suspensión",
            "perfil": "Montaje de Perfil CR",
            "hilo": "Tendido de Hilo de Contacto",
            "calculation": "Avance planificado lineal entre inicio y fin, ponderado con los mismos pesos del KPI real",
            "coverage_rule": "Solo soportes en túneles con cronograma; estaciones y Tramo 2 se reportan fuera de cobertura",
        },
        "excluded_scope": EXCLUDED_SCOPE,
        "source_issues": issues,
        "stages": stages,
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Uso: python extract_schedule.py input.xlsx output.json")
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    payload = extract(source)
    with target.open("w", encoding="utf-8", newline="\n") as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2)
        stream.write("\n")
    print(json.dumps({
        "output": str(target),
        "stages": {name: len(data["tunnels"]) for name, data in payload["stages"].items()},
        "issues": payload["source_issues"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
